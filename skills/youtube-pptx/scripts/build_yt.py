"""
build_yt.py — Builder de slides para canal YouTube · Valeria Yashan PM & Strategy
Paleta oficial PMI · Calibri · 1920×1080

Uso:
    from scripts.build_yt import COVER, SDIV, BULLETS, QUOTE, DETAIL, COMP, STATS, THANKS
    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)
    COVER(prs, "Título del video", "Subtítulo", "Ep. 01")
    BULLETS(prs, "En este video vas a aprender", ["Punto 1", "Punto 2", "Punto 3"])
    THANKS(prs, "¡Gracias por ver!", "Suscribite y activá la campanita 🔔")
    prs.save("output.pptx")
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ─── Paleta PMI ──────────────────────────────────────────────────────────────
NAVY       = RGBColor(0,   48,  135)
BLUE       = RGBColor(0,  102,  204)
CYAN       = RGBColor(0,  174,  239)
ORANGE     = RGBColor(247, 148,  29)
LIGHT_GRAY = RGBColor(240, 244, 248)
DARK_TEXT  = RGBColor(26,   43,  74)
MUTED      = RGBColor(74,   96, 128)
WHITE      = RGBColor(255, 255, 255)

# ─── Helpers internos ────────────────────────────────────────────────────────

def _blank_slide(prs):
    blank_layout = prs.slide_layouts[6]  # completamente en blanco
    return prs.slides.add_slide(blank_layout)


def _rect(slide, left, top, width, height, fill_color, line=False):
    shape = slide.shapes.add_shape(
        1,  # MSO_SHAPE_TYPE RECTANGLE
        left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if not line:
        shape.line.fill.background()
    return shape


def _textbox(slide, left, top, width, height, text, font_name="Calibri",
             font_size=18, bold=False, color=DARK_TEXT, align=PP_ALIGN.LEFT,
             word_wrap=True):
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf = txb.text_frame
    tf.word_wrap = word_wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return txb


def _cyan_line(slide, prs):
    W = prs.slide_width
    _rect(slide, 0, Inches(0.38), W, Pt(4), CYAN)


def _footer(slide, prs):
    W, H = prs.slide_width, prs.slide_height
    _textbox(
        slide,
        Inches(0.5), H - Inches(0.35),
        W - Inches(1), Inches(0.3),
        "@ValeriaYashanPM · PM & Strategy",
        font_size=10, color=MUTED, word_wrap=False
    )


def _title_bar(slide, prs, title_text):
    W = prs.slide_width
    _textbox(
        slide,
        Inches(0.5), Inches(0.55),
        W - Inches(1), Inches(0.6),
        title_text,
        font_size=28, bold=True, color=DARK_TEXT
    )


# ─── Helpers públicos ────────────────────────────────────────────────────────

def COVER(prs, titulo, subtitulo, episodio=""):
    """Portada oscura navy. Siempre primer slide."""
    slide = _blank_slide(prs)
    W, H = prs.slide_width, prs.slide_height

    # Fondo Navy
    _rect(slide, 0, 0, W, H, NAVY)

    # Línea Cyan decorativa (top)
    _rect(slide, 0, Inches(0.3), W, Pt(5), CYAN)

    # Etiqueta episodio (si la hay)
    if episodio:
        _textbox(
            slide,
            Inches(0.6), Inches(1.0),
            Inches(4), Inches(0.4),
            episodio.upper(),
            font_size=14, bold=True, color=ORANGE
        )

    # Título principal
    _textbox(
        slide,
        Inches(0.6), Inches(1.6),
        W - Inches(1.2), Inches(2.0),
        titulo,
        font_size=40, bold=True, color=WHITE
    )

    # Subtítulo
    _textbox(
        slide,
        Inches(0.6), Inches(3.8),
        W - Inches(1.2), Inches(0.8),
        subtitulo,
        font_size=22, color=CYAN
    )

    # Línea Cyan decorativa (bottom)
    _rect(slide, 0, H - Inches(0.5), W, Pt(5), CYAN)

    # Branding bottom-right
    _textbox(
        slide,
        W - Inches(5), H - Inches(0.45),
        Inches(4.5), Inches(0.35),
        "@ValeriaYashanPM · PM & Strategy",
        font_size=11, color=MUTED, align=PP_ALIGN.RIGHT, word_wrap=False
    )
    return slide


def SDIV(prs, seccion, num_seccion=""):
    """Separador de sección. Entre bloques de contenido."""
    slide = _blank_slide(prs)
    W, H = prs.slide_width, prs.slide_height

    # Fondo Navy
    _rect(slide, 0, 0, W, H, NAVY)

    # Número de sección grande (si lo hay)
    if num_seccion:
        _textbox(
            slide,
            Inches(0.6), Inches(1.5),
            Inches(2), Inches(1.5),
            str(num_seccion),
            font_size=60, bold=True, color=CYAN
        )

    # Nombre de sección
    left_offset = Inches(2.8) if num_seccion else Inches(0.6)
    _textbox(
        slide,
        left_offset, Inches(2.2),
        W - left_offset - Inches(0.6), Inches(1.2),
        seccion,
        font_size=30, bold=True, color=WHITE
    )

    # Línea Cyan
    _rect(slide, 0, H - Inches(0.6), W, Pt(5), CYAN)
    return slide


def BULLETS(prs, titulo, puntos):
    """
    Lista con bullets PMI. Máx 5 puntos.
    puntos: lista de strings.
    """
    assert len(puntos) <= 5, "Máximo 5 bullets por slide"
    slide = _blank_slide(prs)
    W, H = prs.slide_width, prs.slide_height

    # Fondo claro
    _rect(slide, 0, 0, W, H, LIGHT_GRAY)
    _cyan_line(slide, prs)
    _title_bar(slide, prs, titulo)
    _footer(slide, prs)

    # Bullets
    txb = slide.shapes.add_textbox(Inches(0.6), Inches(1.35), W - Inches(1.2), H - Inches(2.2))
    tf = txb.text_frame
    tf.word_wrap = True

    for i, punto in enumerate(puntos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"▸  {punto}"
        run = p.runs[0]
        run.font.name = "Calibri"
        run.font.size = Pt(18)
        run.font.color.rgb = DARK_TEXT
        p.space_after = Pt(10)
    return slide


def QUOTE(prs, cita, autor=""):
    """Cita de impacto. Reforzar un concepto clave."""
    slide = _blank_slide(prs)
    W, H = prs.slide_width, prs.slide_height

    # Fondo Navy
    _rect(slide, 0, 0, W, H, NAVY)
    _rect(slide, 0, Inches(0.3), W, Pt(5), CYAN)

    # Comillas decorativas
    _textbox(
        slide, Inches(0.5), Inches(0.8),
        Inches(1.5), Inches(1.5),
        "\u201c",
        font_size=80, bold=True, color=CYAN
    )

    # Texto de la cita
    _textbox(
        slide,
        Inches(0.6), Inches(1.8),
        W - Inches(1.2), Inches(3.0),
        cita,
        font_size=24, color=WHITE, align=PP_ALIGN.CENTER
    )

    # Autor
    if autor:
        _textbox(
            slide,
            Inches(0.6), H - Inches(1.5),
            W - Inches(1.2), Inches(0.5),
            f"— {autor}",
            font_size=16, color=CYAN, align=PP_ALIGN.CENTER
        )
    return slide


def DETAIL(prs, titulo, cuerpo):
    """Texto libre. Explicación extendida."""
    slide = _blank_slide(prs)
    W, H = prs.slide_width, prs.slide_height

    _rect(slide, 0, 0, W, H, LIGHT_GRAY)
    _cyan_line(slide, prs)
    _title_bar(slide, prs, titulo)
    _footer(slide, prs)

    _textbox(
        slide,
        Inches(0.6), Inches(1.35),
        W - Inches(1.2), H - Inches(2.2),
        cuerpo,
        font_size=18, color=DARK_TEXT
    )
    return slide


def COMP(prs, titulo, izq_titulo, izq_contenido, der_titulo, der_contenido):
    """
    2 columnas comparativas. Contraste / pros-cons.
    izq_contenido / der_contenido: string (puede tener \\n para saltos).
    """
    slide = _blank_slide(prs)
    W, H = prs.slide_width, prs.slide_height
    col_w = W / 2 - Inches(0.7)

    _rect(slide, 0, 0, W, H, LIGHT_GRAY)
    _cyan_line(slide, prs)
    _title_bar(slide, prs, titulo)
    _footer(slide, prs)

    # Columna izquierda
    _rect(slide, Inches(0.5), Inches(1.35), col_w, H - Inches(2.1), NAVY)
    _textbox(slide, Inches(0.7), Inches(1.5), col_w - Inches(0.4), Inches(0.5),
             izq_titulo, font_size=16, bold=True, color=CYAN)
    _textbox(slide, Inches(0.7), Inches(2.1), col_w - Inches(0.4), H - Inches(3.0),
             izq_contenido, font_size=16, color=WHITE)

    # Columna derecha
    right_left = W / 2 + Inches(0.2)
    _rect(slide, right_left, Inches(1.35), col_w, H - Inches(2.1), BLUE)
    _textbox(slide, right_left + Inches(0.2), Inches(1.5), col_w - Inches(0.4), Inches(0.5),
             der_titulo, font_size=16, bold=True, color=CYAN)
    _textbox(slide, right_left + Inches(0.2), Inches(2.1), col_w - Inches(0.4), H - Inches(3.0),
             der_contenido, font_size=16, color=WHITE)
    return slide


def STATS(prs, titulo, stats):
    """
    Tarjetas con números. stats: lista de dicts {"numero": "...", "label": "..."}.
    2–4 stats recomendados.
    """
    assert 2 <= len(stats) <= 4, "Entre 2 y 4 stats"
    slide = _blank_slide(prs)
    W, H = prs.slide_width, prs.slide_height

    _rect(slide, 0, 0, W, H, LIGHT_GRAY)
    _cyan_line(slide, prs)
    _title_bar(slide, prs, titulo)
    _footer(slide, prs)

    n = len(stats)
    card_w = (W - Inches(1.0)) / n - Inches(0.2)
    card_colors = [NAVY, BLUE, NAVY, BLUE]

    for i, stat in enumerate(stats):
        left = Inches(0.5) + i * (card_w + Inches(0.2))
        top = Inches(1.5)
        card_h = H - Inches(2.4)

        _rect(slide, left, top, card_w, card_h, card_colors[i % 2])

        # Número grande
        _textbox(
            slide, left + Inches(0.1), top + Inches(0.3),
            card_w - Inches(0.2), Inches(1.2),
            stat["numero"],
            font_size=36, bold=True, color=ORANGE, align=PP_ALIGN.CENTER
        )
        # Label
        _textbox(
            slide, left + Inches(0.1), top + Inches(1.6),
            card_w - Inches(0.2), card_h - Inches(1.8),
            stat["label"],
            font_size=14, color=MUTED, align=PP_ALIGN.CENTER
        )
    return slide


def THANKS(prs, mensaje, cta):
    """Cierre con CTA. Siempre último slide."""
    slide = _blank_slide(prs)
    W, H = prs.slide_width, prs.slide_height

    # Fondo Navy
    _rect(slide, 0, 0, W, H, NAVY)
    _rect(slide, 0, Inches(0.3), W, Pt(5), CYAN)

    # Mensaje principal
    _textbox(
        slide,
        Inches(0.6), Inches(1.8),
        W - Inches(1.2), Inches(1.5),
        mensaje,
        font_size=36, bold=True, color=WHITE, align=PP_ALIGN.CENTER
    )

    # CTA (call to action)
    _textbox(
        slide,
        Inches(0.6), Inches(3.5),
        W - Inches(1.2), Inches(1.0),
        cta,
        font_size=22, color=CYAN, align=PP_ALIGN.CENTER
    )

    # Branding
    _textbox(
        slide,
        Inches(0.6), H - Inches(0.9),
        W - Inches(1.2), Inches(0.4),
        "@ValeriaYashanPM · PM & Strategy",
        font_size=13, color=MUTED, align=PP_ALIGN.CENTER, word_wrap=False
    )

    _rect(slide, 0, H - Inches(0.5), W, Pt(5), CYAN)
    return slide
